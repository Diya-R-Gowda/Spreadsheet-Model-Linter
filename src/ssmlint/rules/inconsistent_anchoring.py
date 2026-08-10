"""inconsistent-anchoring — the fourth and final named Tier 0 rule.

Flags a formula inside a block whose `$`-anchoring on a specific
reference axis (row or column) breaks from the anchoring the rest of the
block establishes for that same logical position — e.g. every member
correctly holds a shared growth-rate cell fixed as `$B$1`, but one
member has it drift to `B1` or `$B1`. Right now the deviant formula may
still compute correctly (the literal address is unchanged — only the
anchoring metadata differs), but it will silently break the moment
anyone copies it to another cell, since only the anchored axis stays
fixed.

WHERE THIS SURFACES (investigated before writing any rule logic, same
discipline as 6a/6b/6c)
------------------------------------------------------------------------
r1c1.py already encodes anchoring directly in its output — an absolute
axis becomes an unbracketed position, a relative axis a bracketed
offset — so a formula with different anchoring on the same logical
reference already normalizes to a different R1C1 string than its
block-mates and cannot cluster into the same block by exact-match
alone. Confirmed with three constructed scenarios (a block holding
`$B$1` fixed, one member drifting): dropping BOTH `$` signs at once
(`$B$1` -> `B1`) landed in `non_conforming` (`different_formula`) — too
large a string change to read as "similar." Dropping just ONE axis's
`$` (`$B$1` -> `$B1`, or `$B$1` -> `B$1`) landed in `near_misses`
instead, scoring 88-91 — comfortably above the near-miss threshold, but
still a rapidfuzz-score artifact, not a guarantee, per the same lesson
learned in range_boundary.py's investigation. So this rule inspects
BOTH `non_conforming` (`kind == "different_formula"`) and `near_misses`
on every block, and independently verifies structurally that the
deviation is a pure anchoring change (same real target cell, different
`$` treatment) rather than trusting whichever bucket blocks.py sorted
it into.

HOW THIS DIFFERS FROM range_boundary.py (verified, not assumed)
---------------------------------------------------------------------
range_boundary detects a shifted BOUNDARY: the deviant formula's range
argument resolves to a genuinely DIFFERENT cell/range than its
block-mates (an off-by-one). inconsistent-anchoring detects the SAME
logical reference position rendered with different anchoring — the
resolved absolute target is IDENTICAL, only the `$` treatment differs.
These two conditions are structurally mutually exclusive by
construction: `_find_anchoring_deltas` below only ever fires when a
differing R1C1 token resolves to the *same* absolute position on every
differing axis (via the candidate cell's own origin); range_boundary's
own check only ever fires when a range endpoint resolves to a
*different* position (shifted by exactly one). A single deviation can
satisfy at most one of the two conditions. Verified directly, not just
argued: constructed both deviation shapes and ran both rule classes
against each — each fires only on its own shape, never on the other's
(see test_inconsistent_anchoring.py).

DETECTION: PER-AXIS, VIA R1C1 STRING STRUCTURE, NOT RE-PARSED FORMULAS
------------------------------------------------------------------------
Every `R..C..` token in both the block's pattern and the candidate's
pattern is located (regex over the fully-regular R1C1 output grammar —
same technique as range_boundary.py, no need to re-parse the underlying
formula). Token counts must match and everything *outside* the tokens
must be byte-for-byte identical (same function, same nesting, same
operators) or this rule stays silent — that's a different, unrelated
deviation. For each token pair that differs, each axis (row, col) is
compared independently, matching r1c1.py's own per-axis anchoring
model:
  - same raw text -> no issue on this axis.
  - different text, same bracket style (both relative or both
    absolute) -> a genuinely different reference, not an anchoring
    issue; the whole candidate is abandoned (not this rule's territory).
  - different text, different bracket style -> resolve the relative
    side to an absolute position using the *candidate cell's own real
    address* as origin (the one piece of concrete, unambiguous position
    data available — a block's own relative offsets have no single
    fixed origin, since they're shared across every member). If the
    resolved positions agree, this is a confirmed pure anchoring
    delta; if they don't, the whole candidate is abandoned (different
    cell, out of scope for this rule).

SEVERITY: THE SAME TWO-LEVEL SCHEME AS EVERY PRIOR RULE, REUSED
------------------------------------------------------------------------
"high" when the deviant cell borders blocks on both sides sharing the
identical pattern (unambiguous single interrupted pattern); "medium"
otherwise. Reused deliberately rather than inventing a new threshold —
same structural argument applies regardless of what kind of deviation
is being evaluated, as already established in literal_in_block.py and
range_boundary.py.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl.utils import column_index_from_string, get_column_letter

from ..ast_nodes import CellRef
from ..blocks import Block, SheetBlocks
from ..tokenizer import parse_cell_addr_parts
from .base import Issue, Rule

RULE_ID = "inconsistent-anchoring"

_RC_TOKEN_RE = re.compile(r"R(\[-?\d+\]|\d+)C(\[-?\d+\]|\d+)")


@dataclass(frozen=True)
class _AnchoringDelta:
    axis: str  # "row" | "col"
    expected_style: str  # "absolute" | "relative"
    actual_style: str  # "absolute" | "relative"
    reference: str  # A1-style address of the real target cell, e.g. "B1"
    expected_notation: str  # e.g. "$B$1"
    actual_notation: str  # e.g. "$B1"


def _origin(address: str) -> tuple[int, int]:
    _sheet, plain = address.split("!", 1)
    _col_abs, col, _row_abs, row = parse_cell_addr_parts(plain)
    return row, column_index_from_string(col)


def _parse_token(token: str) -> tuple[str, str]:
    match = _RC_TOKEN_RE.match(token)
    assert match is not None  # guaranteed by _RC_TOKEN_RE having already matched
    return match.group(1), match.group(2)


def _style(raw: str) -> str:
    return "relative" if raw.startswith("[") else "absolute"


def _resolve(raw: str, origin_val: int) -> int:
    return origin_val + int(raw.strip("[]")) if raw.startswith("[") else int(raw)


def _find_anchoring_deltas(
    block_pattern: str, candidate_pattern: str, origin_row: int, origin_col: int
) -> list[_AnchoringDelta] | None:
    block_tokens = list(_RC_TOKEN_RE.finditer(block_pattern))
    candidate_tokens = list(_RC_TOKEN_RE.finditer(candidate_pattern))
    if not block_tokens or len(block_tokens) != len(candidate_tokens):
        return None

    block_skeleton = _RC_TOKEN_RE.sub("\0", block_pattern)
    candidate_skeleton = _RC_TOKEN_RE.sub("\0", candidate_pattern)
    if block_skeleton != candidate_skeleton:
        return None  # something outside the reference tokens also differs

    deltas: list[_AnchoringDelta] = []
    for block_match, candidate_match in zip(block_tokens, candidate_tokens):
        if block_match.group(0) == candidate_match.group(0):
            continue

        b_row_raw, b_col_raw = _parse_token(block_match.group(0))
        c_row_raw, c_col_raw = _parse_token(candidate_match.group(0))
        c_abs_row = _resolve(c_row_raw, origin_row)
        c_abs_col = _resolve(c_col_raw, origin_col)
        reference = f"{get_column_letter(c_abs_col)}{c_abs_row}"
        expected_notation = CellRef(
            col=get_column_letter(c_abs_col), row=c_abs_row,
            col_abs=_style(b_col_raw) == "absolute", row_abs=_style(b_row_raw) == "absolute",
        ).to_a1()
        actual_notation = CellRef(
            col=get_column_letter(c_abs_col), row=c_abs_row,
            col_abs=_style(c_col_raw) == "absolute", row_abs=_style(c_row_raw) == "absolute",
        ).to_a1()

        for axis, b_raw, c_raw, origin_val in (
            ("row", b_row_raw, c_row_raw, origin_row),
            ("col", b_col_raw, c_col_raw, origin_col),
        ):
            if b_raw == c_raw:
                continue
            b_axis_style, c_axis_style = _style(b_raw), _style(c_raw)
            if b_axis_style == c_axis_style:
                return None  # same anchoring style, different value -> a different reference
            if _resolve(b_raw, origin_val) != _resolve(c_raw, origin_val):
                return None  # doesn't resolve to the same cell -> not a pure anchoring issue
            deltas.append(
                _AnchoringDelta(
                    axis=axis, expected_style=b_axis_style, actual_style=c_axis_style,
                    reference=reference, expected_notation=expected_notation,
                    actual_notation=actual_notation,
                )
            )

    return deltas if deltas else None


def _anchoring_touchpoints(sheet_blocks: list[SheetBlocks]) -> dict[str, list[tuple[Block, list[_AnchoringDelta]]]]:
    touchpoints: dict[str, list[tuple[Block, list[_AnchoringDelta]]]] = {}
    for sheet in sheet_blocks:
        for block in sheet.blocks:
            candidates: list[tuple[str, str | None]] = [
                (nc.cell, nc.formula_r1c1)
                for nc in block.non_conforming
                if nc.kind == "different_formula"
            ]
            candidates += [(nm.cell, nm.formula_r1c1) for nm in block.near_misses]
            for cell, pattern in candidates:
                if pattern is None:
                    continue
                origin_row, origin_col = _origin(cell)
                deltas = _find_anchoring_deltas(block.pattern, pattern, origin_row, origin_col)
                if deltas:
                    touchpoints.setdefault(cell, []).append((block, deltas))
    return touchpoints


def _severity(bordering: list[tuple[Block, list[_AnchoringDelta]]]) -> str:
    if len(bordering) == 2 and bordering[0][0].pattern == bordering[1][0].pattern:
        return "high"
    return "medium"


def _explanation(cell: str, bordering: list[tuple[Block, list[_AnchoringDelta]]]) -> str:
    block0, deltas0 = bordering[0]
    spans = " and ".join(f"{block.sheet}!{block.span}" for block, _d in bordering)
    axis_parts = [
        f"the {d.axis} axis of its reference to {d.reference} is {d.actual_style} "
        f"({d.actual_notation}) instead of {d.expected_style} ({d.expected_notation})"
        for d in deltas0
    ]
    return (
        f"{cell}: {'; '.join(axis_parts)} — every other member of {spans} "
        f"(following {block0.pattern}) anchors this reference consistently."
    )


def _suggested_fix(bordering: list[tuple[Block, list[_AnchoringDelta]]]) -> str:
    _block0, deltas0 = bordering[0]
    notations = sorted({d.expected_notation for d in deltas0})
    return (
        f"This block's established anchoring for this reference is {', '.join(notations)} — if this "
        f"cell were meant to follow the pattern, its anchoring would match that instead of drifting, "
        f"so the reference stays fixed (or moves) exactly like the rest of the block if ever copied."
    )


class InconsistentAnchoringRule(Rule):
    rule_id = RULE_ID

    def evaluate(self, sheet_blocks: list[SheetBlocks], graph=None) -> list[Issue]:
        issues: list[Issue] = []
        for cell, bordering in _anchoring_touchpoints(sheet_blocks).items():
            issues.append(
                Issue(
                    cell=cell,
                    severity=_severity(bordering),
                    rule_id=RULE_ID,
                    explanation=_explanation(cell, bordering),
                    suggested_fix=_suggested_fix(bordering),
                )
            )
        return issues
