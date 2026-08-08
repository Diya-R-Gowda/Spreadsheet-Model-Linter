"""Workbook parsing spine.

Loads an .xlsx/.xlsm workbook and extracts, per cell, the raw formula
string (never the cached computed value), literal values, number formats,
and merged-range membership. Also resolves workbook- and sheet-scoped
named ranges, and flags formulas that reference another sheet.

This module deliberately does not tokenize formulas or build an AST —
that is a later pipeline stage. A formula this module can't represent
cleanly as a plain string (array formulas, LET/LAMBDA, structured table
references) is still captured best-effort, but is also logged to the
`skipped` list so later stages know it needs special handling.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}

# Matches a sheet-qualified reference prefix, e.g. `Sheet2!` or `'My Sheet'!`.
_SHEET_REF_RE = re.compile(r"(?:'(?P<quoted>[^']+)'|(?P<bare>[A-Za-z_][A-Za-z0-9_.]*))!")
_LET_LAMBDA_RE = re.compile(r"\b(LET|LAMBDA)\s*\(", re.IGNORECASE)
_STRUCTURED_REF_RE = re.compile(r"[A-Za-z_][A-Za-z0-9_.]*\[[^\[\]]+\]")


@dataclass
class SkippedItem:
    address: str
    reason: str

    def to_dict(self) -> dict[str, str]:
        return {"address": self.address, "reason": self.reason}


@dataclass
class CellRecord:
    address: str  # sheet-qualified, e.g. "Sheet1!C14"
    formula: str | None
    value: Any
    number_format: str
    merged: bool
    cross_sheet_refs: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "address": self.address,
            "formula": self.formula,
            "value": self.value,
            "number_format": self.number_format,
            "merged": self.merged,
            "cross_sheet_refs": self.cross_sheet_refs,
        }


@dataclass
class SheetRecord:
    name: str
    cells: list[CellRecord] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {"name": self.name, "cells": [c.to_dict() for c in self.cells]}


@dataclass
class NamedRangeRecord:
    name: str
    refers_to: str
    scope: str  # "workbook" or the sheet name it's local to

    def to_dict(self) -> dict[str, str]:
        return {"name": self.name, "refers_to": self.refers_to, "scope": self.scope}


@dataclass
class ParsedWorkbook:
    workbook: str
    sheets: list[SheetRecord] = field(default_factory=list)
    named_ranges: list[NamedRangeRecord] = field(default_factory=list)
    skipped: list[SkippedItem] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "workbook": self.workbook,
            "sheets": [s.to_dict() for s in self.sheets],
            "named_ranges": [n.to_dict() for n in self.named_ranges],
            "skipped": [s.to_dict() for s in self.skipped],
        }


def _find_cross_sheet_refs(formula: str, current_sheet: str) -> list[str]:
    refs: list[str] = []
    for match in _SHEET_REF_RE.finditer(formula):
        sheet_name = match.group("quoted") or match.group("bare")
        if sheet_name and sheet_name != current_sheet and sheet_name not in refs:
            refs.append(sheet_name)
    return refs


def _extract_formula_text(raw: Any, address: str, skipped: list[SkippedItem]) -> str:
    """Return a best-effort raw formula string, logging shapes we don't fully handle yet."""
    if isinstance(raw, ArrayFormula):
        skipped.append(SkippedItem(address, "array formula (not tokenized in this step)"))
        text = raw.text
    elif isinstance(raw, str):
        text = raw
    else:
        text = str(raw)

    if _LET_LAMBDA_RE.search(text):
        skipped.append(SkippedItem(address, "LET/LAMBDA formula (not tokenized in this step)"))
    if _STRUCTURED_REF_RE.search(text):
        skipped.append(SkippedItem(address, "structured table reference (not tokenized in this step)"))

    return text


def _extract_named_ranges(wb: Workbook) -> list[NamedRangeRecord]:
    records: list[NamedRangeRecord] = []
    for name, defined_name in wb.defined_names.items():
        records.append(NamedRangeRecord(name=name, refers_to=defined_name.value or "", scope="workbook"))
    for ws in wb.worksheets:
        for name, defined_name in ws.defined_names.items():
            records.append(NamedRangeRecord(name=name, refers_to=defined_name.value or "", scope=ws.title))
    return records


def _parse_sheet(ws: Worksheet, skipped: list[SkippedItem]) -> SheetRecord:
    sheet_name = ws.title
    merged_ranges = list(ws.merged_cells.ranges)
    sheet_record = SheetRecord(name=sheet_name)

    for row in ws.iter_rows():
        for cell in row:
            address = f"{sheet_name}!{cell.coordinate}"
            try:
                is_merged = any(cell.coordinate in mr for mr in merged_ranges)

                formula: str | None = None
                value: Any = None
                cross_refs: list[str] = []

                if cell.data_type == "f":
                    formula = _extract_formula_text(cell.value, address, skipped)
                    cross_refs = _find_cross_sheet_refs(formula, sheet_name)
                else:
                    value = cell.value

                sheet_record.cells.append(
                    CellRecord(
                        address=address,
                        formula=formula,
                        value=value,
                        number_format=cell.number_format,
                        merged=is_merged,
                        cross_sheet_refs=cross_refs,
                    )
                )
            except Exception as exc:  # deliberate catch-all: never let one bad cell abort the parse
                skipped.append(SkippedItem(address, f"unparseable cell: {exc!r}"))

    return sheet_record


def parse_workbook(path: str | Path) -> ParsedWorkbook:
    """Load a workbook and extract its parsing-spine representation.

    Raises ValueError for unsupported file extensions. Per-cell failures
    are caught and recorded in the result's `skipped` list rather than
    raised, so a single malformed cell never aborts the whole parse.
    """
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {path.suffix!r} (expected .xlsx or .xlsm)")

    wb = openpyxl.load_workbook(path, data_only=False)

    result = ParsedWorkbook(workbook=path.name)
    result.named_ranges = _extract_named_ranges(wb)

    for ws in wb.worksheets:
        result.sheets.append(_parse_sheet(ws, result.skipped))

    return result
